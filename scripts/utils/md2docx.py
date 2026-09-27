#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Markdown -> Word (.docx) 转换器
支持: 标题层级(#~####)、段落、加粗(**text**)、行内代码(`code`)、
      引用上标([1] -> ^[1])、Markdown 表格、水平线、列表。
中文正文使用宋体，标题使用黑体，西文 Times New Roman。
"""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


CN_FONT = "宋体"
CN_HEAD = "黑体"
EN_FONT = "Times New Roman"


def set_run_font(run, cn=CN_FONT, en=EN_FONT, size=None, bold=None, italic=None):
    run.font.name = en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), en)
    rFonts.set(qn('w:hAnsi'), en)
    rFonts.set(qn('w:eastAsia'), cn)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic


def add_formatted_text(paragraph, text, base_size=11.5, base_bold=False, cn=CN_FONT):
    """解析 **bold**、`code`、[n] 上标 并写入 paragraph"""
    # tokenizer: split on bold, inline-code, citation
    pattern = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[\d+(?:[,\-–]\d+)*\])')
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            seg = text[pos:m.start()]
            if seg:
                r = paragraph.add_run(seg)
                set_run_font(r, cn=cn, size=base_size, bold=base_bold)
        tok = m.group(0)
        if tok.startswith('**'):
            r = paragraph.add_run(tok[2:-2])
            set_run_font(r, cn=cn, size=base_size, bold=True)
        elif tok.startswith('`'):
            r = paragraph.add_run(tok[1:-1])
            set_run_font(r, cn=cn, en="Consolas", size=base_size - 1)
        else:  # citation [1]
            r = paragraph.add_run(tok)
            set_run_font(r, cn=cn, size=base_size)
            r.font.superscript = True
        pos = m.end()
    if pos < len(text):
        seg = text[pos:]
        if seg:
            r = paragraph.add_run(seg)
            set_run_font(r, cn=cn, size=base_size, bold=base_bold)


def set_cell_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), '808080')
        borders.append(el)
    tcPr.append(borders)


def shade_cell(cell, color="EFEFEF"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def convert(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = Document()

    # 页面设置 A4
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.17)
    sec.right_margin = Cm(3.17)

    # 正文默认样式
    normal = doc.styles['Normal']
    normal.font.name = EN_FONT
    normal.font.size = Pt(11.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # 空行
        if not stripped:
            i += 1
            continue

        # 水平线
        if re.match(r'^---+$', stripped):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:color'), '999999')
            pbdr.append(bottom)
            pPr.append(pbdr)
            i += 1
            continue

        # 表格
        if stripped.startswith('|'):
            tbl_lines = []
            while i < n and lines[i].strip().startswith('|'):
                tbl_lines.append(lines[i].strip())
                i += 1
            # 去掉分隔行
            rows = []
            for tl in tbl_lines:
                if re.match(r'^\|[\s\-:|]+\|$', tl):
                    continue
                cells = [c.strip() for c in tl.strip('|').split('|')]
                rows.append(cells)
            if not rows:
                continue
            ncol = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=ncol)
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(rows):
                for ci in range(ncol):
                    cell = table.cell(ri, ci)
                    cell.text = ''
                    para = cell.paragraphs[0]
                    val = row[ci] if ci < len(row) else ''
                    val = val.replace('<br>', ' ')
                    add_formatted_text(para, val, base_size=9.5,
                                       base_bold=(ri == 0),
                                       cn=(CN_HEAD if ri == 0 else CN_FONT))
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    para.paragraph_format.space_before = Pt(2)
                    para.paragraph_format.space_after = Pt(2)
                    set_cell_borders(cell)
                    if ri == 0:
                        shade_cell(cell)
            doc.add_paragraph()
            continue

        # 标题
        m = re.match(r'^(#{1,6})\s+(.*)$', stripped)
        if m:
            level = len(m.group(1))
            title = m.group(2).strip()
            if level == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(18)
                add_formatted_text(p, title, base_size=18, base_bold=True, cn=CN_HEAD)
            elif level == 2:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(16)
                p.paragraph_format.space_after = Pt(10)
                add_formatted_text(p, title, base_size=15, base_bold=True, cn=CN_HEAD)
            elif level == 3:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(13)
                p.paragraph_format.space_after = Pt(8)
                add_formatted_text(p, title, base_size=13, base_bold=True, cn=CN_HEAD)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(6)
                add_formatted_text(p, title, base_size=12, base_bold=True, cn=CN_HEAD)
            i += 1
            continue

        # 无序列表
        m = re.match(r'^[-*]\s+(.*)$', stripped)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            add_formatted_text(p, m.group(1), base_size=11.5)
            i += 1
            continue

        # 有序列表
        m = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            add_formatted_text(p, m.group(2), base_size=11.5)
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.space_after = Pt(6)
        pf.first_line_indent = Pt(23)  # 中文首行缩进 2 字符
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_formatted_text(p, stripped, base_size=11.5)
        i += 1

    doc.save(docx_path)
    print(f"OK -> {docx_path}")


if __name__ == '__main__':
    convert(sys.argv[1], sys.argv[2])
